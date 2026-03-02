"""
Pyffice PowerPoint Handler
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pathlib import Path
from typing import List, Optional


def create() -> Presentation:
    """Create new PowerPoint presentation."""
    return Presentation()


def add_slide(prs: Presentation, layout_index: int = 1) -> object:
    """Add slide to presentation."""
    slide = prs.slides.add_slide(prs.slide_layouts[layout_index])
    return slide


def add_title(slide: object, title: str, font_size: int = 32) -> None:
    """Add title to slide."""
    title_shape = slide.shapes.title
    title_shape.text = title
    title_shape.text_frame.paragraphs[0].font.size = Pt(font_size)


def add_text(slide: object, text: str, left: float = 1, top: float = 1.5, width: float = 8, height: float = 4) -> None:
    """Add text box to slide."""
    textbox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = textbox.text_frame
    tf.text = text


def add_image(slide: object, image_path: str, left: float = 1, top: float = 1.5, width: Optional[float] = None) -> None:
    """Add image to slide."""
    if width:
        slide.shapes.add_picture(image_path, Inches(left), Inches(top), width=Inches(width))
    else:
        slide.shapes.add_picture(image_path, Inches(left), Inches(top))


def save(prs: Presentation, filepath: str) -> None:
    """Save presentation to file."""
    prs.save(filepath)


def open(filepath: str) -> Presentation:
    """Open existing PowerPoint file."""
    return Presentation(filepath)


def get_slides(prs: Presentation) -> List:
    """Get list of slides."""
    return list(prs.slides)
